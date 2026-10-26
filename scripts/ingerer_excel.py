"""Charge les fichiers Excel de saisie dans la table raw.raw_saisie.

Chaque ligne garde sa provenance (fichier, hash, onglet, n° de ligne, date de
chargement, run_id) pour la traçabilité. Les valeurs sont stockées en texte :
le typage et les contrôles sont faits en SQL (dbt), donc portables Snowflake.
Un fichier déjà chargé (même hash) n'est pas rechargé.
"""
import argparse
import hashlib
import os
import uuid
from datetime import datetime, timezone
from pathlib import Path

import duckdb
import yaml
from openpyxl import load_workbook

RACINE = Path(__file__).resolve().parents[1]
DB_DEFAUT = os.environ.get("BC_DUCKDB_PATH", str(RACINE / "bilan_carbone.duckdb"))

DDL = """
create schema if not exists raw;
create table if not exists raw.raw_saisie (
    run_id           varchar,
    fichier_source   varchar,
    hash_fichier     varchar,
    onglet           varchar,
    numero_ligne     integer,
    charge_le        timestamp,
    campagne         integer,
    domaine          varchar,
    territoire       varchar,
    libelle          varchar,
    quantite         varchar,
    unite            varchar,
    data_owner       varchar,
    commentaire      varchar
);
"""
COLONNES = ["domaine", "territoire", "libelle", "quantite", "unite", "data_owner", "commentaire"]


def hash_fichier(chemin: Path) -> str:
    return hashlib.sha256(chemin.read_bytes()).hexdigest()


def normaliser(texte) -> str:
    return str(texte).strip().lower() if texte is not None else ""


def lire_fichier(chemin: Path, mapping: dict) -> list[dict]:
    wb = load_workbook(chemin, read_only=True, data_only=True)
    onglet = mapping["onglet_saisie"] if mapping["onglet_saisie"] in wb.sheetnames else wb.sheetnames[0]
    lignes = list(wb[onglet].iter_rows(values_only=True))
    if not lignes:
        return []

    entetes = [normaliser(h) for h in lignes[0]]
    index = {}
    for cible, variantes in mapping["colonnes"].items():
        for v in variantes:
            if normaliser(v) in entetes:
                index[cible] = entetes.index(normaliser(v))
                break
    manquantes = [c for c in mapping["obligatoires"] if c not in index]
    if manquantes:
        raise ValueError(f"{chemin.name} : colonnes obligatoires introuvables {manquantes}")

    resultat = []
    for num, valeurs in enumerate(lignes[1:], start=2):
        if all(v is None or str(v).strip() == "" for v in valeurs):
            continue
        enreg = {"onglet": onglet, "numero_ligne": num}
        for c in COLONNES:
            v = valeurs[index[c]] if c in index and index[c] < len(valeurs) else None
            enreg[c] = None if v is None else str(v).strip()
        resultat.append(enreg)
    return resultat


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--dossier", default=str(RACINE / "data" / "entree"))
    parser.add_argument("--campagne", type=int, default=2026)
    parser.add_argument("--db", default=DB_DEFAUT)
    args = parser.parse_args()

    mapping = yaml.safe_load((RACINE / "config" / "mapping_excel.yml").read_text(encoding="utf-8"))
    run_id = uuid.uuid4().hex[:12]
    maintenant = datetime.now(timezone.utc).replace(tzinfo=None)

    con = duckdb.connect(args.db)
    con.execute(DDL)
    deja = {r[0] for r in con.execute("select distinct hash_fichier from raw.raw_saisie").fetchall()}

    total = 0
    for chemin in sorted(Path(args.dossier).glob("*.xlsx")):
        if chemin.name.startswith("~$"):
            continue
        h = hash_fichier(chemin)
        if h in deja:
            print(f"ignoré (déjà chargé) : {chemin.name}")
            continue
        lignes = lire_fichier(chemin, mapping)
        con.executemany(
            "insert into raw.raw_saisie values (?,?,?,?,?,?,?,?,?,?,?,?,?,?)",
            [
                [run_id, chemin.name, h, l["onglet"], l["numero_ligne"], maintenant, args.campagne]
                + [l[c] for c in COLONNES]
                for l in lignes
            ],
        )
        total += len(lignes)
        print(f"chargé : {chemin.name} ({len(lignes)} lignes)")
    con.close()
    print(f"run {run_id} : {total} lignes ingérées")


if __name__ == "__main__":
    main()
